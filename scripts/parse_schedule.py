#!/usr/bin/env python3
"""Parse the KFU timetable workbook into a compact JSON schedule.

The workbook uses merged cells extensively. A lesson may start several columns
to the left of the requested group, so reading only the group's column loses
common lectures. This parser resolves every merged range that intersects the
selected group column.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


GROUP_RE = re.compile(r"^\s*\d{2}-\d{3}(?:\s*\([^)]*\))*\s*$")
TIME_RE = re.compile(r"(?P<sh>\d{1,2})[.:](?P<sm>\d{2})\s*[-–—]\s*(?P<eh>\d{1,2})[.:](?P<em>\d{2})")
DAY_IDS = {
    "понедельник": "monday",
    "вторник": "tuesday",
    "среда": "wednesday",
    "четверг": "thursday",
    "пятница": "friday",
    "суббота": "saturday",
}
DAY_NAMES = {
    "monday": "Понедельник",
    "tuesday": "Вторник",
    "wednesday": "Среда",
    "thursday": "Четверг",
    "friday": "Пятница",
    "saturday": "Суббота",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def canonical_group(value: str) -> str:
    return re.sub(r"\s+", "", value).casefold()


def group_slug(value: str) -> str:
    """Return a stable, URL-safe filename for a group."""
    readable = re.sub(r"[^0-9a-z]+", "-", value.casefold()).strip("-")
    readable = readable[:48].rstrip("-") or "group"
    digest = hashlib.sha1(canonical_group(value).encode("utf-8")).hexdigest()[:8]
    return f"{readable}-{digest}"


def normalize_time(value: str) -> tuple[str, str, str] | None:
    match = TIME_RE.search(clean_text(value))
    if not match:
        return None
    start = f"{int(match.group('sh')):02d}:{match.group('sm')}"
    end = f"{int(match.group('eh')):02d}:{match.group('em')}"
    return f"{start}-{end}", start, end


@dataclass(frozen=True)
class GroupLocation:
    name: str
    header_row: int
    column: int
    course: str
    program: str
    day_column: int
    time_column: int


class TimetableParser:
    def __init__(self, workbook_path: Path, sheet_name: str | None = None):
        self.workbook_path = workbook_path
        self.workbook = load_workbook(workbook_path, data_only=True, read_only=False)
        self.sheet = self.workbook[sheet_name] if sheet_name else self.workbook.active
        self._merged_ranges = list(self.sheet.merged_cells.ranges)
        self._merged_by_cell: dict[tuple[int, int], Any] = {}
        for merged_range in self._merged_ranges:
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    self._merged_by_cell[(row, col)] = merged_range

    def value_at(self, row: int, column: int) -> str:
        merged_range = self._merged_by_cell.get((row, column))
        if merged_range:
            row, column = merged_range.min_row, merged_range.min_col
        return clean_text(self.sheet.cell(row, column).value)

    def direct_value_at(self, row: int, column: int) -> str:
        """Return only the physical cell value, without repeating a merged anchor."""
        return clean_text(self.sheet.cell(row, column).value)

    def _course_marker(self, header_row: int, group_column: int) -> tuple[int, str]:
        for row in range(header_row, max(0, header_row - 4), -1):
            for col in range(group_column, 0, -1):
                value = self.value_at(row, col)
                if re.search(r"\b\d+\s*курс\b|магистратур", value, re.IGNORECASE):
                    merged = self._merged_by_cell.get((row, col))
                    marker_col = merged.min_col if merged else col
                    return marker_col, value
        raise ValueError(f"Не найден блок курса рядом с {get_column_letter(group_column)}{header_row}")

    def find_groups(self) -> list[GroupLocation]:
        groups: list[GroupLocation] = []
        seen: set[tuple[int, int]] = set()
        for row in self.sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                value = clean_text(cell.value)
                if not value or not GROUP_RE.match(value):
                    continue
                key = (cell.row, cell.column)
                if key in seen:
                    continue
                seen.add(key)
                day_col, course = self._course_marker(cell.row, cell.column)
                program = self.value_at(cell.row - 1, cell.column)
                if program == course:
                    program = ""
                groups.append(
                    GroupLocation(
                        name=value,
                        header_row=cell.row,
                        column=cell.column,
                        course=course,
                        program=program,
                        day_column=day_col,
                        time_column=day_col + 1,
                    )
                )
        return groups

    def select_group(self, requested: str) -> GroupLocation:
        wanted = canonical_group(requested)
        exact = [group for group in self.find_groups() if canonical_group(group.name) == wanted]
        if len(exact) == 1:
            return exact[0]
        partial = [group for group in self.find_groups() if wanted in canonical_group(group.name)]
        if len(partial) == 1:
            return partial[0]
        if not exact and not partial:
            raise ValueError(f"Группа {requested!r} не найдена. Используйте --list-groups.")
        names = ", ".join(group.name for group in (exact or partial))
        raise ValueError(f"Название группы неоднозначно: {names}")

    def _schedule_end_row(self, group: GroupLocation) -> int:
        start = group.header_row + 1
        saw_saturday = False
        for row in range(start, min(self.sheet.max_row, start + 140) + 1):
            day_value = self.direct_value_at(row, group.day_column).casefold()
            if day_value == "суббота":
                saw_saturday = True
            if saw_saturday and day_value in {"вск", "воскресенье"}:
                return row - 1
        return min(self.sheet.max_row, start + 100)

    def _lesson_values(self, group: GroupLocation, first_row: int, last_row: int) -> list[tuple[str, str]]:
        found: list[tuple[str, str]] = []
        seen_sources: set[str] = set()
        for row in range(first_row, last_row + 1):
            merged_range = self._merged_by_cell.get((row, group.column))
            if merged_range:
                source = str(merged_range)
                value = clean_text(self.sheet.cell(merged_range.min_row, merged_range.min_col).value)
            else:
                source = f"{get_column_letter(group.column)}{row}"
                value = clean_text(self.sheet.cell(row, group.column).value)
            if value and source not in seen_sources:
                seen_sources.add(source)
                found.append((value, source))
        return found

    def parse_group(self, group: GroupLocation) -> dict[str, Any]:
        end_row = self._schedule_end_row(group)
        days = {day_id: {"id": day_id, "name": DAY_NAMES[day_id], "lessons": []} for day_id in DAY_NAMES}
        current_day: str | None = None
        row = group.header_row + 1

        while row <= end_row:
            raw_day = self.direct_value_at(row, group.day_column).casefold()
            if raw_day in DAY_IDS:
                current_day = DAY_IDS[raw_day]

            time_info = normalize_time(self.direct_value_at(row, group.time_column))
            if not time_info or not current_day:
                row += 1
                continue

            next_row = row + 1
            while next_row <= end_row:
                if normalize_time(self.direct_value_at(next_row, group.time_column)):
                    break
                next_day = self.direct_value_at(next_row, group.day_column).casefold()
                if next_day in DAY_IDS or next_day in {"вт", "ср", "чт", "пт", "сб", "вск"}:
                    break
                next_row += 1

            time_value, start_time, end_time = time_info
            raw_lessons = self._lesson_values(group, row, next_row - 1)
            for raw_text, source in raw_lessons:
                for part in split_scheduled_entries(raw_text):
                    lesson = parse_lesson(part)
                    lesson.update(
                        {
                            "time": time_value,
                            "start_time": start_time,
                            "end_time": end_time,
                            "source_cell": source,
                        }
                    )
                    days[current_day]["lessons"].append(lesson)
            row = next_row

        for day in days.values():
            day["lessons"].sort(key=lambda lesson: (lesson["start_time"], lesson["weeks"], lesson["subject"]))

        return {
            "source": {
                "file": self.workbook_path.name,
                "sheet": self.sheet.title,
                "group": group.name,
                "course": group.course,
                "program": group.program,
                "group_column": get_column_letter(group.column),
            },
            "week_count": 18,
            "days": list(days.values()),
        }

    def parse(self, requested_group: str) -> dict[str, Any]:
        return self.parse_group(self.select_group(requested_group))


def split_scheduled_entries(text: str) -> list[str]:
    # Split only when the next semicolon-delimited part starts with another
    # week rule. Teacher/room lists separated by semicolons stay together.
    parts = re.split(
        r";\s*(?=\(\s*(?:(?:н|ч)\s*/\s*н\s*)?(?:с\s*)?\d[^)]*нед)",
        clean_text(text),
        flags=re.IGNORECASE,
    )
    return [part.strip(" ;") for part in parts if part.strip(" ;")]


def parse_week_rule(text: str) -> tuple[dict[str, Any], str]:
    rule_match = re.match(r"^\s*\(([^)]*нед[^)]*)\)\s*", text, flags=re.IGNORECASE)
    if not rule_match:
        return {
            "raw": "",
            "parity": None,
            "start_week": 1,
            "end_week": 18,
            "weeks": list(range(1, 19)),
        }, text.strip()

    raw_rule = clean_text(rule_match.group(1)).replace("–", "-").replace("—", "-")
    lowered = raw_rule.casefold()
    parity = "odd" if re.search(r"н\s*/\s*н", lowered) else "even" if re.search(r"ч\s*/\s*н", lowered) else None
    weeks: set[int] = set()
    for start, end in re.findall(r"(\d{1,2})\s*-\s*(\d{1,2})", lowered):
        first, last = int(start), int(end)
        if first <= last:
            weeks.update(range(first, last + 1))
    if not weeks:
        singles = [int(value) for value in re.findall(r"\b(\d{1,2})\b", lowered)]
        weeks.update(value for value in singles if 1 <= value <= 18)
    if not weeks:
        weeks.update(range(1, 19))
    if parity == "odd":
        weeks = {week for week in weeks if week % 2 == 1}
    elif parity == "even":
        weeks = {week for week in weeks if week % 2 == 0}

    ordered_weeks = sorted(week for week in weeks if 1 <= week <= 18)
    rule = {
        "raw": raw_rule,
        "parity": parity,
        "start_week": min(ordered_weeks) if ordered_weeks else None,
        "end_week": max(ordered_weeks) if ordered_weeks else None,
        "weeks": ordered_weeks,
    }
    return rule, text[rule_match.end() :].strip()


def parse_lesson(text: str) -> dict[str, Any]:
    rule, body = parse_week_rule(clean_text(text))
    # Пометка «-д» в конце записи означает дополнительную пару, а не часть аудитории.
    additional = bool(re.search(r"\s*[-–—]\s*д\.?\s*$", body, flags=re.IGNORECASE))
    if additional:
        body = re.sub(r"\s*[-–—]\s*д\.?\s*$", "", body, flags=re.IGNORECASE).rstrip(" ,.;")
    room = ""
    before_room = body
    room_match = re.search(r"\bауд\.?\s*(.+)$", body, flags=re.IGNORECASE)
    if room_match:
        room = clean_text(room_match.group(1)).rstrip(".")
        before_room = body[: room_match.start()].rstrip(" ,.-")

    teacher = ""
    subject = before_room
    teacher_match = re.search(r"[А-ЯЁ][А-Яа-яё-]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.?", before_room)
    if teacher_match:
        subject = before_room[: teacher_match.start()].rstrip(" ,.-")
        teacher = before_room[teacher_match.start() :].strip(" ,.-")

    lowered = body.casefold()
    lesson_type = "lesson"
    if "эор" in lowered:
        lesson_type = "online"
    elif re.search(r"\bлаб", lowered):
        lesson_type = "laboratory"
    elif re.search(r"\bлек", lowered):
        lesson_type = "lecture"
    elif "практик" in lowered:
        lesson_type = "practice"
    elif "физическ" in lowered:
        lesson_type = "physical_education"

    return {
        "subject": subject or body,
        "teacher": teacher,
        "room": room,
        "additional": additional,
        "type": lesson_type,
        "text": body,
        "rule": {key: value for key, value in rule.items() if key != "weeks"},
        "weeks": rule["weeks"],
    }


def group_rows(groups: Iterable[GroupLocation]) -> list[dict[str, Any]]:
    return [
        {
            "group": group.name,
            "course": group.course,
            "program": group.program,
            "cell": f"{get_column_letter(group.column)}{group.header_row}",
        }
        for group in groups
    ]


def build_cli() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Парсер расписания КФУ из Excel в JSON")
    parser.add_argument("input", type=Path, help="Путь к файлу .xlsx")
    parser.add_argument("--sheet", help="Название листа; по умолчанию активный лист")
    parser.add_argument("--group", help='Группа, например "09-551 (1)"')
    parser.add_argument("--list-groups", action="store_true", help="Вывести найденные группы")
    parser.add_argument("--all-groups", action="store_true", help="Создать отдельный JSON для каждой найденной группы")
    parser.add_argument("--output-dir", type=Path, help="Каталог для --all-groups; рядом создаётся groups.json")
    parser.add_argument("--output", "-o", type=Path, help="Файл результата; без него JSON выводится в консоль")
    parser.add_argument("--compact", action="store_true", help="JSON без отступов")
    return parser


def main() -> int:
    args = build_cli().parse_args()
    if not args.input.exists():
        print(f"Файл не найден: {args.input}", file=sys.stderr)
        return 2
    try:
        parser = TimetableParser(args.input, args.sheet)
        if args.list_groups:
            result: Any = group_rows(parser.find_groups())
        elif args.all_groups:
            if not args.output_dir:
                print("Для --all-groups укажите --output-dir.", file=sys.stderr)
                return 2

            groups = parser.find_groups()
            args.output_dir.mkdir(parents=True, exist_ok=True)
            index_groups: list[dict[str, Any]] = []

            for group in groups:
                file_name = f"{group_slug(group.name)}.json"
                schedule = parser.parse_group(group)
                serialized_schedule = json.dumps(
                    schedule,
                    ensure_ascii=False,
                    indent=None if args.compact else 2,
                )
                (args.output_dir / file_name).write_text(serialized_schedule + "\n", encoding="utf-8")
                index_groups.append(
                    {
                        "id": group_slug(group.name),
                        "name": group.name,
                        "course": group.course,
                        "program": group.program,
                        "file": f"groups/{file_name}",
                    }
                )

            index_result = {
                "version": 1,
                "source": {"file": args.input.name, "sheet": parser.sheet.title},
                "groups": index_groups,
            }
            index_path = args.output_dir.parent / "groups.json"
            index_path.write_text(
                json.dumps(index_result, ensure_ascii=False, indent=None if args.compact else 2) + "\n",
                encoding="utf-8",
            )
            print(f"Создано расписаний: {len(index_groups)}")
            print(index_path)
            return 0
        else:
            if not args.group:
                print("Укажите --group или используйте --list-groups.", file=sys.stderr)
                return 2
            result = parser.parse(args.group)
    except (KeyError, ValueError) as error:
        print(str(error), file=sys.stderr)
        return 2

    serialized = json.dumps(result, ensure_ascii=False, indent=None if args.compact else 2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(serialized + "\n", encoding="utf-8")
        print(args.output)
    else:
        print(serialized)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
